# -*- coding: utf-8 -*-
"""回读五个结果文件，核对结构、行数、总量与列对齐。"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pathlib import Path
import numpy as np
import openpyxl

RESULTS = Path(__file__).resolve().parents[1] / "results"

for name in ["result1.xlsx", "result2.xlsx", "result3.xlsx", "result4-2.xlsx", "result4-3.xlsx"]:
    p = RESULTS / name
    wb = openpyxl.load_workbook(p, data_only=True)
    print(f"\n########## {name}  sheets={wb.sheetnames}  大小={p.stat().st_size/1024:.0f} KB")
    for ws in wb.worksheets:
        print(f"  -- {ws.title}: {ws.max_row} 行 x {ws.max_column} 列")
        if ws.max_column > 10:
            head = [ws.cell(row=1, column=c).value for c in (1, 2, 145, 146, 147)]
            first = [ws.cell(row=2, column=c).value for c in (1, 2, 145, 146, 147)]
            last = [ws.cell(row=ws.max_row, column=c).value for c in (1, 2, 145, 146, 147)]
            print("     header:", head)
            print("     row2  :", first)
            print("     last  :", last)
        else:
            for r in list(range(1, 4)) + list(range(max(4, ws.max_row - 1), ws.max_row + 1)):
                print(f"     r{r:<4}", [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)])

print("\n\n########## 结构与总量核对 ##########")
wb = openpyxl.load_workbook(RESULTS / "result2.xlsx", data_only=True)
ws = wb["计划购电量"]
vals = np.array([[ws.cell(row=r, column=c).value for c in range(2, 146)] for r in range(2, 336)], dtype=float)
tot = np.array([ws.cell(row=r, column=146).value for r in range(2, 336)], dtype=float)
print("result2 计划购电量: 形状", vals.shape, " 无空值:", not np.isnan(vals).any())
print("  行内和 vs 全天购电量 最大偏差:", np.abs(vals.sum(axis=1) - tot).max())
print("  首日第一个区间(应为 0:00-0:10) =", vals[0, 0], " 末个区间 =", vals[0, -1])

ws = wb["充放电量"]
print("result2 充放电量 行数:", ws.max_row, "(应为 1 + 334*6 =", 1 + 334 * 6, ")")
days = [ws.cell(row=2 + 6 * i, column=1).value for i in range(334)]
print("  首/末日期:", days[0], days[-1], " 日期列无空:", all(d is not None for d in days))
soc0 = np.array([ws.cell(row=2 + 6 * i, column=6).value for i in range(334)], dtype=float)
soc24 = np.array([ws.cell(row=3 + 6 * i, column=6).value for i in range(334)], dtype=float)
print("  0:00 储电量 min/max:", soc0.min(), soc0.max(), " 24:00 储电量 min/max:", soc24.min(), soc24.max())
print("  首尾是否相等:", np.abs(soc0 - soc24).max())
chg = np.array([[ws.cell(row=2 + 6 * i + k, column=3).value for k in range(6)] for i in range(334)], dtype=float)
dis = np.array([[ws.cell(row=2 + 6 * i + k, column=4).value for k in range(6)] for i in range(334)], dtype=float)
print("  全充电量 %.2f kWh  全放电量 %.2f kWh" % (chg.sum(), dis.sum()))

ws = wb["紧急购电量"]
print("result2 紧急购电量 行数:", ws.max_row, "(空表应为 2)")
n_run = sum(1 for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=2).value)
amt = sum(ws.cell(row=r, column=3).value or 0 for r in range(2, ws.max_row + 1))
print("  紧急购电区间条数:", n_run, " 电量合计 %.2f kWh" % amt)

ws = openpyxl.load_workbook(RESULTS / "result3.xlsx", data_only=True)["调整购电量"]
print("result3 调整购电量 行数:", ws.max_row, " 首个全天购电量:", ws.cell(row=2, column=146).value)
ws = openpyxl.load_workbook(RESULTS / "result1.xlsx", data_only=True)["计划购电量"]
xs = np.array([ws.cell(row=r, column=2).value for r in range(2, 146)], dtype=float)
print("result1 购电量合计 %.4f kWh" % xs.sum())
