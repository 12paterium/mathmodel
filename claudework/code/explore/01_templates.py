# -*- coding: utf-8 -*-
"""精查模板表头与行数、附件1/4 时间标签口径。"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pathlib import Path
import openpyxl

DATA = Path(__file__).resolve().parent.parent / 'data'

def head_tail(path, sheet, row_lo, row_hi, cols=(1,2,3,4,143,144,145,146,147)):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    print(f"\n=== {path.name} :: {sheet}  max_row={ws.max_row} max_col={ws.max_column}")
    r = row_lo
    while r <= row_hi:
        vals = []
        for c in cols:
            if c <= ws.max_column:
                v = ws.cell(row=r, column=c).value
                vals.append(f"c{c}={v}")
        print(f"  r{r:<4}", " | ".join(vals))
        r += 1

print("############ 模板 result1 计划购电量 ############")
head_tail(DATA/'templates'/'result1.xlsx', '计划购电量', 1, 3, cols=(1,2))
head_tail(DATA/'templates'/'result1.xlsx', '计划购电量', 142, 145, cols=(1,2))

print("\n############ 模板 result2 计划购电量 ############")
head_tail(DATA/'templates'/'result2.xlsx', '计划购电量', 1, 2, cols=(1,2,3,144,145,146,147))
head_tail(DATA/'templates'/'result2.xlsx', '计划购电量', 333, 335, cols=(1,2,3,144,145,146,147))

print("\n############ 模板 result3 表头 ############")
head_tail(DATA/'templates'/'result3.xlsx', '计划购电量', 1, 1, cols=(1,2,3,143,144,145,146,147))
head_tail(DATA/'templates'/'result3.xlsx', '调整购电量', 1, 2, cols=(1,2,3,144,145,146,147))
head_tail(DATA/'templates'/'result3.xlsx', '调整购电量', 333, 335, cols=(1,2,3,144,145,146,147))
for s in ['充放电量','紧急购电量']:
    wb = openpyxl.load_workbook(DATA/'templates'/'result3.xlsx', data_only=True)
    ws = wb[s]
    print(f"\n=== result3 :: {s} max_row={ws.max_row} max_col={ws.max_column}")
    for r in list(range(1,4)) + list(range(ws.max_row-2, ws.max_row+1)):
        print(f"  r{r:<4}", [ws.cell(row=r, column=c).value for c in range(1, ws.max_column+1)])

print("\n############ 附件1 尾部 ############")
wb = openpyxl.load_workbook(DATA/'附件1.xlsx', data_only=True)
ws = wb['Sheet1']
for r in [140, 141, 142, 143, 144, 145]:
    print(f"  r{r}", [ws.cell(row=r, column=c).value for c in range(1,5)])
print("  r1 header:", [ws.cell(row=1, column=c).value for c in range(1,5)])

print("\n############ 附件4 尾部表头 ############")
wb = openpyxl.load_workbook(DATA/'附件4.xlsx', data_only=True)
ws = wb['Sheet1']
print("  last col ctx:", [(c, ws.cell(row=1, column=c).value) for c in range(141, 146)])
print("  r2 tail      :", [ws.cell(row=2, column=c).value for c in range(140, 146)])
print("  A366..A366  :", ws.cell(row=366, column=1).value, " max_row", ws.max_row)

print("\n############ 附件3 尾部 ############")
wb = openpyxl.load_workbook(DATA/'附件3.xlsx', data_only=True)
ws = wb['Sheet1']
print("  header cols:", [(c, ws.cell(row=1, column=c).value) for c in range(1, 27)])
for r in [1457, 1458, 1459, 1460, 1461]:
    print(f"  r{r}", [ws.cell(row=r, column=c).value for c in range(1, 5)])
