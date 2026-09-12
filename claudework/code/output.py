# -*- coding: utf-8 -*-
"""把四问结果写进附件5 的模板。

附件5 的“计划购电量”表头是从 0:10-0:20 开始写的，比当天实际区间整体错开
一格（当天第一个区间是 0:00-0:10）。模板结构不能改，所以按行号映射：
第 2 行放第 0 个区间（0:00-0:10），依次到第 145 行放第 143 个区间（23:50-24:00）。
"""
from pathlib import Path

import numpy as np
import openpyxl

from data_io import DATA, SOC_INIT

RESULTS = Path(__file__).resolve().parents[1] / "results"
BLOCKS = ["0:00-4:00", "4:00-8:00", "8:00-12:00",
          "12:00-16:00", "16:00-20:00", "20:00-24:00"]


def _load(template):
    return openpyxl.load_workbook(DATA / "templates" / template)


def _as_date(value):
    """numpy 日期转成 Excel 认的 Python date。"""
    return np.datetime64(value, "D").astype(object)


def _block_sums(values):
    return values.reshape(6, 24).sum(axis=1)


def write_result1(ins, dispatch):
    """问题1：144 段购电量 + 6 个时段的充放电量与首尾储电量。"""
    wb = _load("result1.xlsx")
    sheet = wb["计划购电量"]
    for i in range(144):
        sheet.cell(row=2 + i, column=2).value = round(float(dispatch.purchase[i]), 4)

    sheet = wb["充放电量"]
    charge = _block_sums(dispatch.charge)
    discharge = _block_sums(dispatch.discharge)
    for k in range(6):
        sheet.cell(row=2 + k, column=2).value = round(float(charge[k]), 4)
        sheet.cell(row=2 + k, column=3).value = round(float(discharge[k]), 4)
    sheet.cell(row=2, column=5).value = round(SOC_INIT, 4)
    sheet.cell(row=3, column=5).value = round(float(dispatch.soc[-1]), 4)
    wb.save(RESULTS / "result1.xlsx")


def write_wide_sheet(sheet, days, values, price, start_row=2):
    """把 (天数, 144) 的购电量写进“日期\\时间”形式的宽表，并补全天两列。

    全天购电量、全天购电费由四舍五入后的分段值累加，保证和表内数字对得上。
    """
    values = np.round(np.asarray(values), 4)
    price = np.asarray(price)
    for i, day in enumerate(days):
        row = start_row + i
        for t in range(144):
            sheet.cell(row=row, column=2 + t).value = float(values[i, t])
        sheet.cell(row=row, column=146).value = round(float(values[i].sum()), 4)
        sheet.cell(row=row, column=147).value = round(float((values[i] * price[i]).sum()), 2)


def write_storage_sheet(sheet, days, charge, discharge, soc_start, soc_end):
    """按天写 6 个时段的充放电量、0:00 与 24:00 的储电量。"""
    charge = np.asarray(charge)
    discharge = np.asarray(discharge)
    for i, day in enumerate(days):
        top = 2 + 6 * i
        sheet.cell(row=top, column=1).value = _as_date(day)
        charge_blocks = _block_sums(charge[i])
        discharge_blocks = _block_sums(discharge[i])
        for k in range(6):
            row = top + k
            sheet.cell(row=row, column=2).value = BLOCKS[k]
            sheet.cell(row=row, column=3).value = round(float(charge_blocks[k]), 4)
            sheet.cell(row=row, column=4).value = round(float(discharge_blocks[k]), 4)
        sheet.cell(row=top, column=5).value = "0:00"
        sheet.cell(row=top, column=6).value = round(float(soc_start[i]), 4)
        sheet.cell(row=top + 1, column=5).value = "24:00"
        sheet.cell(row=top + 1, column=6).value = round(float(soc_end[i]), 4)


def write_emergency_sheet(sheet, days, runs_per_day):
    """每天一行日期，下面跟着该天的连续紧急购电区间与电量。"""
    row = 2
    for day, runs in zip(days, runs_per_day):
        sheet.cell(row=row, column=1).value = _as_date(day)
        for start_time, end_time, kwh in runs:
            row += 1
            sheet.cell(row=row, column=2).value = f"{start_time}-{end_time}"
            sheet.cell(row=row, column=3).value = round(kwh, 4)
        row += 1


def write_result2(ins, records):
    """问题2：计划购电量、充放电量、紧急购电量。"""
    days = [ins.dates[r["day"]] for r in records[31:]]
    price = [ins.a1_price for _ in records[31:]]
    purchase = [r["purchase"] for r in records[31:]]
    charge = [r["charge"] for r in records[31:]]
    discharge = [r["discharge"] for r in records[31:]]
    soc_start = [r["soc_start"] for r in records[31:]]
    soc_end = [r["soc_end"] for r in records[31:]]

    wb = _load("result2.xlsx")
    write_wide_sheet(wb["计划购电量"], days, purchase, price)
    write_storage_sheet(wb["充放电量"], days, charge, discharge, soc_start, soc_end)
    write_emergency_sheet(wb["紧急购电量"], days, [r["runs"] for r in records[31:]])
    wb.save(RESULTS / "result2.xlsx")


def write_result3(ins, records):
    """问题3：0:00 计划购电量、调整购电量、最终充放电量、紧急购电量。"""
    span = records[31:]
    days = [ins.dates[r["day"]] for r in span]
    price = [ins.a1_price for _ in span]

    wb = _load("result3.xlsx")
    write_wide_sheet(wb["计划购电量"], days, [r["base_purchase"] for r in span], price)
    write_wide_sheet(wb["调整购电量"], days, [r["purchase"] for r in span], price)
    write_storage_sheet(wb["充放电量"], days, [r["charge"] for r in span],
                        [r["discharge"] for r in span],
                        [r["soc_start"] for r in span], [r["soc_end"] for r in span])
    write_emergency_sheet(wb["紧急购电量"], days, [r["runs"] for r in span])
    wb.save(RESULTS / "result3.xlsx")


def write_result4(ins, records42, records43):
    """问题4：波动电价下的问题2 与问题3。"""
    span2, span3 = records42[31:], records43[31:]
    days2 = [ins.dates[r["day"]] for r in span2]
    days3 = [ins.dates[r["day"]] for r in span3]
    price2 = [ins.price[r["day"]] for r in span2]
    price3 = [ins.price[r["day"]] for r in span3]

    wb = _load("result4-2.xlsx")
    write_wide_sheet(wb["计划购电量"], days2, [r["purchase"] for r in span2], price2)
    write_storage_sheet(wb["充放电量"], days2, [r["charge"] for r in span2],
                        [r["discharge"] for r in span2],
                        [r["soc_start"] for r in span2], [r["soc_end"] for r in span2])
    write_emergency_sheet(wb["紧急购电量"], days2, [r["runs"] for r in span2])
    wb.save(RESULTS / "result4-2.xlsx")

    wb = _load("result4-3.xlsx")
    write_wide_sheet(wb["计划购电量"], days3, [r["base_purchase"] for r in span3], price3)
    write_wide_sheet(wb["调整购电量"], days3, [r["purchase"] for r in span3], price3)
    write_storage_sheet(wb["充放电量"], days3, [r["charge"] for r in span3],
                        [r["discharge"] for r in span3],
                        [r["soc_start"] for r in span3], [r["soc_end"] for r in span3])
    write_emergency_sheet(wb["紧急购电量"], days3, [r["runs"] for r in span3])
    wb.save(RESULTS / "result4-3.xlsx")
