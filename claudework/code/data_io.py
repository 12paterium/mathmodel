# -*- coding: utf-8 -*-
"""读取附件1-4，输出规整的 numpy 数组。

约定（功率 -> 电量）：所有功率乘 1/6 小时，得到 10 分钟区间的电量 kWh。

时间口径：
- 附件1 的“时间”列与附件2/4 的表头（0:10 ... 0:00+1）是区间的**结束时刻**，
  共 144 个区间，对应 0:00-0:10 ... 23:50-24:00。
- 附件3 的一行是一次预报，列“预报h小时”指从发布时刻起第 h 个小时
  （即区间 (T+h-1, T+h]）的光伏功率预报；使用时在该小时的 6 个十分钟段内取常值。
"""
from pathlib import Path
from dataclasses import dataclass

import numpy as np
import openpyxl

DATA = Path(__file__).resolve().parents[1] / "data"

DT = 1.0 / 6.0            # 小时
N = 144                   # 每天区间数
FLOW_MAX = 5000.0 * DT    # 单区间最大充/放电量，kWh
SOC_MIN, SOC_MAX = 1200.0, 10800.0
SOC_INIT = 6000.0
ETA = 0.9                 # 题面“充放电效率为 90%”

RELEASE_HOURS = (0, 6, 12, 18)
DISPLAY_DATES = ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21")


def _sheet_grid(sheet):
    """把 '日期\\时间' 形式的表读成 (dates[365], values[365,144])。"""
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    dates = np.array([np.datetime64(r[0].date()) for r in rows])
    values = np.array([[float(v) for v in r[1:1 + N]] for r in rows])
    return dates, values


@dataclass
class Inputs:
    dates: np.ndarray     # (365,) 日期
    a1_price: np.ndarray  # (144,) 元/kWh
    a1_load: np.ndarray   # (144,) 区间电量 kWh
    a1_pv: np.ndarray     # (144,) 区间电量 kWh
    load: np.ndarray      # (365,144) 小区负载 区间电量 kWh
    pv: np.ndarray        # (365,144) 实际光伏 区间电量 kWh
    price: np.ndarray     # (365,144) 实时电价 元/kWh
    pv_fc: np.ndarray     # (365,4,24) 光伏预报（每小时常值，已换算为 kWh）


def load_inputs():
    wb1 = openpyxl.load_workbook(DATA / "附件1.xlsx", read_only=True, data_only=True)
    rows = list(wb1.worksheets[0].iter_rows(min_row=2, values_only=True))
    a1_price = np.array([r[1] for r in rows], dtype=float)
    a1_load = np.array([r[2] for r in rows], dtype=float) * DT
    a1_pv = np.array([r[3] for r in rows], dtype=float) * DT

    wb2 = openpyxl.load_workbook(DATA / "附件2.xlsx", read_only=True, data_only=True)
    dates, load = _sheet_grid(wb2["小区负载"])
    _, pv = _sheet_grid(wb2["光伏发电实际功率"])

    wb4 = openpyxl.load_workbook(DATA / "附件4.xlsx", read_only=True, data_only=True)
    _, price = _sheet_grid(wb4.worksheets[0])

    # 附件3：日期只写在每天第一行，读取时向下填充
    wb3 = openpyxl.load_workbook(DATA / "附件3.xlsx", read_only=True, data_only=True)
    date_index = {d: i for i, d in enumerate(dates)}
    pv_fc = np.zeros((len(dates), len(RELEASE_HOURS), 24))
    current = None
    for row in wb3.worksheets[0].iter_rows(min_row=2, values_only=True):
        if row[0] not in (None, ""):
            text = row[0] if isinstance(row[0], str) else row[0].strftime("%Y-%m-%d")
            y, m, d = (int(v) for v in text.split("-"))
            current = date_index[np.datetime64(f"{y:04d}-{m:02d}-{d:02d}")]
        hour = int(str(row[1]).split(":")[0])
        pv_fc[current, RELEASE_HOURS.index(hour), :] = [float(v) for v in row[2:26]]
    pv_fc *= DT

    return Inputs(dates=dates, a1_price=a1_price, a1_load=a1_load, a1_pv=a1_pv,
                  load=load * DT, pv=pv * DT, price=price, pv_fc=pv_fc)


def interval_labels():
    """144 个区间的 'h:mm-h:mm' 标签，第 0 个是 0:00-0:10。"""
    return [f"{i * 10 // 60}:{i * 10 % 60:02d}-{(i + 1) * 10 // 60}:{(i + 1) * 10 % 60:02d}"
            for i in range(N)]


def clock_label(index):
    """第 index 个区间边界的时刻：0 是 0:00，144 是 24:00。"""
    minutes = index * 10
    return f"{minutes // 60}:{minutes % 60:02d}"


def hour_slice(release_hour, hour_offset):
    """发布时刻 release_hour 的第 hour_offset 个小时预报，对应哪 6 个十分钟区间。

    hour_offset=0 表示第 1 小时（“预报1小时”）。
    """
    start = (release_hour + hour_offset) * 6
    return slice(start, start + 6)


def block_sums(values):
    """把长度 144 的数组压成 6 个 4 小时时段的和。"""
    return values.reshape(6, 24).sum(axis=1)
