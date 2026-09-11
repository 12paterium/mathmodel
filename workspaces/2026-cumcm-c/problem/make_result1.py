# -*- coding: utf-8 -*-
"""问题1收尾：无惩罚项复算校验 + 无储能基线对照 + 生成 result1.xlsx"""
import numpy as np
import openpyxl
from scipy.optimize import linprog
from openpyxl.utils import get_column_letter

# ---------------- 读数据与参数 ----------------
wb = openpyxl.load_workbook('附件/附件1.xlsx', data_only=True)
rows = list(wb['Sheet1'].iter_rows(min_row=2, values_only=True))
price = np.array([float(r[1]) for r in rows])
load  = np.array([float(r[2]) for r in rows])
pv    = np.array([float(r[3]) for r in rows])
n = 144; dt = 1/6; eta = 0.9
E0, Emin, Emax, Pmax = 6000.0, 1200.0, 10800.0, 5000.0

# ---------------- 无惩罚项复算（校验最优值） ----------------
N = 4*n
obj = np.zeros(N); obj[:n] = price*dt
A_ub = np.zeros((n, N)); b_ub = pv - load
A_ub[:, :n] = -np.eye(n); A_ub[:, n:2*n] = np.eye(n); A_ub[:, 2*n:3*n] = -np.eye(n)
A_eq = np.zeros((n+1, N)); b_eq = np.zeros(n+1)
for k in range(n):
    A_eq[k, 3*n+k] = 1.0
    if k > 0: A_eq[k, 3*n+k-1] = -1.0
    A_eq[k, n+k] = -dt*eta; A_eq[k, 2*n+k] = dt/eta
b_eq[0] = E0; A_eq[n, 4*n-1] = 1.0; b_eq[n] = E0
bounds = [(0,None)]*n + [(0,Pmax)]*n + [(0,Pmax)]*n + [(Emin,Emax)]*n
res = linprog(obj, A_ub=A_ub, b_ub=b_ub, A_eq=A_eq, b_eq=b_eq, bounds=bounds, method='highs')
print('无惩罚项复算: status=%d, 全天购电费=%.4f 元' % (res.status, res.fun))

# ---------------- 无储能基线（必须全部买电） ----------------
net_load = np.maximum(0.0, load - pv)
base_cost = float((price*net_load*dt).sum())
base_buy  = float((net_load*dt).sum())
print('无储能基线: 购电量=%.2f kWh, 购电费=%.2f 元' % (base_buy, base_cost))

# ---------------- 读取带惩罚的解（采用） ----------------
d = np.load('p1_solution.npz')
buy_e, ch_e, dis_e = d['buy_e'], d['ch_e'], d['dis_e']
E = d['E']
pure_cost = float((price*buy_e).sum())
print('采用方案: 购电量=%.2f kWh, 购电费=%.2f 元 (较基线省 %.2f 元, %.1f%%)'
      % (buy_e.sum(), pure_cost, base_cost-pure_cost, 100*(base_cost-pure_cost)/base_cost))
print('负载总电量=%.2f kWh, 光伏总发电=%.2f kWh' % ((load*dt).sum(), (pv*dt).sum()))
print('充电总量=%.2f kWh, 放电总量=%.2f kWh' % (ch_e.sum(), dis_e.sum()))

def r4(x):
    v = round(float(x), 4)
    return 0.0 if v == 0 else v

# ---------------- 生成 result1.xlsx ----------------
wb2 = openpyxl.load_workbook('附件/附件5/result1.xlsx')
ws1 = wb2['计划购电量']      # 按位置填充：第 i 行 = 第 i 个10分钟时段 (10(i-1), 10i]
assert ws1.max_row == 145
for i in range(144):
    ws1.cell(row=2+i, column=2).value = r4(buy_e[i])

ws2 = wb2['充放电量']        # 6个4小时大段 + 0:00/24:00储电量
blocks = [(0,240),(240,480),(480,720),(720,960),(960,1200),(1200,1440)]
for j,(a,bb) in enumerate(blocks):
    sl = slice(a//10, bb//10)
    ws2.cell(row=2+j, column=2).value = r4(ch_e[sl].sum())   # 充电量
    ws2.cell(row=2+j, column=3).value = r4(dis_e[sl].sum())  # 放电量
ws2.cell(row=2, column=5).value = r4(E0)     # 0:00 储电量
ws2.cell(row=3, column=5).value = r4(E[-1])  # 24:00 储电量

wb2.save('result1.xlsx')
print('\nresult1.xlsx 已生成')

# ---------------- 打印论文用表1/表2 ----------------
def hm(m):
    h,mm = divmod(m,60); return f'{h}:{mm:02d}' if h<24 else '24:00'
print('\n===== 表1 微网在指定时间段的购电量及全天的购电量和购电费 =====')
for t_end in [610, 730, 850, 970, 1090, 1210]:
    k = t_end//10 - 1
    print('  %s-%s: %.2f kWh' % (hm(t_end-10), hm(t_end), buy_e[k]))
print('  全天购电量: %.2f kWh' % buy_e.sum())
print('  全天购电费: %.2f 元' % pure_cost)
print('\n===== 表2 储能设备在指定时间段的充放电量及0:00和24:00的储电量 =====')
for j,(a,bb) in enumerate(blocks):
    sl = slice(a//10, bb//10)
    print('  %s-%s: 充电量 %.2f  放电量 %.2f' % (hm(a), hm(bb), ch_e[sl].sum(), dis_e[sl].sum()))
print('  0:00储电量: %.2f kWh' % E0)
print('  24:00储电量: %.2f kWh' % E[-1])
