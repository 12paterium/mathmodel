"""Inspect source workbooks and submission templates without modifying them."""

from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
ATTACHMENTS = ROOT / "problem" / "附件"


def preview_workbook(path: Path) -> None:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    print(f"\n{path.relative_to(ROOT)}")
    for sheet in workbook.worksheets:
        print(f"  [{sheet.title}] rows={sheet.max_row}, cols={sheet.max_column}")
        for row in sheet.iter_rows(
            min_row=1,
            max_row=min(sheet.max_row, 6),
            max_col=min(sheet.max_column, 8),
            values_only=True,
        ):
            print("   ", row)


def escaped(value: object) -> str:
    return ascii(value)


def inspect_templates() -> None:
    for path in sorted((ATTACHMENTS / "附件5").glob("*.xlsx")):
        workbook = openpyxl.load_workbook(path, read_only=False, data_only=False)
        print(f"\nTEMPLATE {path.name}")
        for sheet in workbook.worksheets:
            headers = [escaped(sheet.cell(1, col).value) for col in range(1, sheet.max_column + 1)]
            print(f"  SHEET {escaped(sheet.title)} HEADERS {headers}")
            if sheet.max_row <= 30:
                for row in range(2, sheet.max_row + 1):
                    values = [escaped(sheet.cell(row, col).value) for col in range(1, sheet.max_column + 1)]
                    print(f"    ROW {row}: {values}")


def main() -> None:
    for number in range(1, 5):
        preview_workbook(ATTACHMENTS / f"附件{number}.xlsx")
    for path in sorted((ATTACHMENTS / "附件5").glob("*.xlsx")):
        preview_workbook(path)
    inspect_templates()


if __name__ == "__main__":
    main()
