"""Shared data, forecasting, optimization, and execution code for C problem."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable  # noqa: F401  (kept for callers)

import numpy as np
import openpyxl
from scipy.optimize import linprog
from scipy.sparse import lil_matrix


ROOT = Path(__file__).resolve().parents[1]
ATTACHMENTS = ROOT / "problem" / "附件"
DT = 1.0 / 6.0
N_INTERVALS = 144
ETA = 0.9
SOC_MIN = 1200.0
SOC_MAX = 10800.0
SOC_INITIAL = 6000.0
FLOW_MAX = 5000.0 * DT
SUBMIT_START = 31
DISPLAY_DATES = ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21")


@dataclass
class Inputs:
    dates: np.ndarray
    fixed_price: np.ndarray
    fixed_load: np.ndarray
    fixed_pv: np.ndarray
    load: np.ndarray
    pv: np.ndarray
    price: np.ndarray
    pv_forecasts: np.ndarray


@dataclass
class Plan:
    purchase: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    soc: np.ndarray
    curtail: np.ndarray
    objective: float
    balance_residual: float


@dataclass
class Execution:
    charge: np.ndarray
    discharge: np.ndarray
    soc: np.ndarray
    emergency: np.ndarray
    curtail: np.ndarray
    balance_residual: float


def _numeric_matrix(sheet: openpyxl.worksheet.worksheet.Worksheet) -> tuple[np.ndarray, np.ndarray]:
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    dates = np.array([np.datetime64(row[0].date()) for row in rows])
    values = np.asarray([[float(value) for value in row[1:145]] for row in rows], dtype=float)
    return dates, values


def load_inputs() -> Inputs:
    wb1 = openpyxl.load_workbook(ATTACHMENTS / "附件1.xlsx", read_only=True, data_only=True)
    rows1 = list(wb1.worksheets[0].iter_rows(min_row=2, values_only=True))
    fixed_price = np.asarray([float(row[1]) for row in rows1])
    fixed_load = np.asarray([float(row[2]) for row in rows1]) * DT
    fixed_pv = np.asarray([float(row[3]) for row in rows1]) * DT

    wb2 = openpyxl.load_workbook(ATTACHMENTS / "附件2.xlsx", read_only=True, data_only=True)
    dates, load = _numeric_matrix(wb2.worksheets[0])
    dates_pv, pv = _numeric_matrix(wb2.worksheets[1])
    if not np.array_equal(dates, dates_pv):
        raise ValueError("附件2的负载与光伏日期不一致")

    wb4 = openpyxl.load_workbook(ATTACHMENTS / "附件4.xlsx", read_only=True, data_only=True)
    dates_price, price = _numeric_matrix(wb4.worksheets[0])
    if not np.array_equal(dates, dates_price):
        raise ValueError("附件4与附件2日期不一致")

    wb3 = openpyxl.load_workbook(ATTACHMENTS / "附件3.xlsx", read_only=True, data_only=True)
    pv_forecasts = np.zeros((len(dates), 4, 24), dtype=float)
    date_to_index = {str(date): index for index, date in enumerate(dates)}
    active_date = None
    release_to_index = {"0:00": 0, "6:00": 1, "12:00": 2, "18:00": 3}
    for row in wb3.worksheets[0].iter_rows(min_row=2, values_only=True):
        if row[0] not in (None, ""):
            active_date = np.datetime64(datetime.strptime(str(row[0]), "%Y-%m-%d").date())
        release = str(row[1]).strip()
        pv_forecasts[date_to_index[str(active_date)], release_to_index[release], :] = np.asarray(row[2:26], float)

    for name, matrix in {"负载": load, "光伏": pv, "实时电价": price}.items():
        if matrix.shape != (365, N_INTERVALS) or not np.isfinite(matrix).all() or (matrix < 0).any():
            raise ValueError(f"{name}数据未通过形状/非负/有限性检查")
    if len(fixed_price) != N_INTERVALS or pv_forecasts.shape != (365, 4, 24):
        raise ValueError("附件1或附件3形状错误")
    return Inputs(
        dates=dates,
        fixed_price=fixed_price,
        fixed_load=fixed_load,
        fixed_pv=fixed_pv,
        load=load * DT,
        pv=pv * DT,
        price=price,
        pv_forecasts=pv_forecasts * DT,
    )


def _base_forecasts(values: np.ndarray) -> np.ndarray:
    """Return causal [previous day, 7-day mean, same weekday] forecasts."""
    days, intervals = values.shape
    bases = np.full((days, 3, intervals), np.nan)
    for day in range(1, days):
        bases[day, 0] = values[day - 1]
        bases[day, 1] = values[max(0, day - 7):day].mean(axis=0)
        weekday_history = np.arange(day - 7, -1, -7)[:4]
        bases[day, 2] = values[weekday_history].mean(axis=0) if len(weekday_history) else bases[day, 1]
    bases[0] = values[0]
    return bases


def causal_ensemble(values: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    bases = _base_forecasts(values)
    forecasts = np.empty_like(values)
    weights = np.empty((len(values), 3), dtype=float)
    forecasts[0] = values[0]
    weights[0] = 1.0 / 3.0
    for day in range(1, len(values)):
        validation = np.arange(max(1, day - 14), day)
        if len(validation) < 2:
            day_weights = np.full(3, 1.0 / 3.0)
        else:
            errors = np.mean(np.abs(bases[validation] - values[validation, None, :]), axis=(0, 2))
            inverse = 1.0 / np.maximum(errors, 1e-9)
            day_weights = inverse / inverse.sum()
        weights[day] = day_weights
        forecasts[day] = np.tensordot(day_weights, bases[day], axes=(0, 0))
    return forecasts, weights


def select_safety_quantile(load: np.ndarray, pv: np.ndarray, load_fc: np.ndarray, pv_fc: np.ndarray) -> tuple[float, dict[str, float]]:
    residual = (load - pv) - (load_fc - pv_fc)
    candidates = (0.70, 0.75, 0.80, 0.85, 0.90)
    validation_days = range(14, 31)
    losses: dict[str, float] = {}
    for quantile in candidates:
        loss = 0.0
        for day in validation_days:
            margin = np.quantile(residual[max(1, day - 14):day], quantile, axis=0)
            error = residual[day] - margin
            loss += float(np.maximum(-error, 0).sum() + 5.0 * np.maximum(error, 0).sum())
        losses[f"{quantile:.2f}"] = loss
    chosen = min(candidates, key=lambda value: losses[f"{value:.2f}"])
    return chosen, losses


def safety_margin(day: int, residual: np.ndarray, quantile: float, lookback: int = 28) -> np.ndarray:
    history = residual[max(1, day - lookback):day]
    if not len(history):
        return np.zeros(N_INTERVALS)
    return np.maximum(0.0, np.quantile(history, quantile, axis=0))


def forecast_metrics(actual: np.ndarray, predicted: np.ndarray, start: int = SUBMIT_START) -> dict[str, float]:
    error = predicted[start:] - actual[start:]
    return {
        "mae": float(np.mean(np.abs(error))),
        "rmse": float(np.sqrt(np.mean(error**2))),
        "wape": float(np.sum(np.abs(error)) / max(np.sum(np.abs(actual[start:])), 1e-12)),
    }


def expanded_pv_forecast(inputs: Inputs, day: int, release_hour: int, bias_correct: bool = True) -> np.ndarray:
    release_index = release_hour // 6
    hours_left = 24 - release_hour
    raw = np.repeat(inputs.pv_forecasts[day, release_index, :hours_left], 6)
    if not bias_correct or day == 0:
        return raw
    history = np.arange(max(0, day - 14), day)
    errors = []
    start = release_hour * 6
    for past in history:
        observed = inputs.pv[past, start:]
        forecast = np.repeat(inputs.pv_forecasts[past, release_index, :hours_left], 6)
        errors.append(observed - forecast)
    correction = np.median(np.asarray(errors), axis=0) if errors else 0.0
    return np.maximum(0.0, raw + correction)


def plan_dispatch(
    load: np.ndarray,
    pv: np.ndarray,
    decision_price: np.ndarray,
    start_soc: float,
    terminal_value: float,
    terminal_soc: float | None = None,
    old_purchase: np.ndarray | None = None,
    eta: float = ETA,
) -> Plan:
    """Solve a deterministic interval-energy LP; all flow variables are kWh."""
    n = len(load)
    if not (len(pv) == len(decision_price) == n):
        raise ValueError("LP输入长度不一致")
    adjusted = old_purchase is not None
    blocks = 7 if adjusted else 5
    total = blocks * n
    x0, c0, d0, s0, w0 = (index * n for index in range(5))
    up0, down0 = (5 * n, 6 * n)

    objective = np.zeros(total)
    if adjusted:
        objective[up0:up0 + n] = 1.5 * decision_price
        objective[down0:down0 + n] = -0.5 * decision_price
    else:
        objective[x0:x0 + n] = decision_price
    objective[c0:c0 + n] = 1e-8
    objective[d0:d0 + n] = 1e-8
    objective[w0:w0 + n] = 1e-9
    objective[s0 + n - 1] -= terminal_value

    eq_rows = 2 * n + (n if adjusted else 0) + (1 if terminal_soc is not None else 0)
    a_eq = lil_matrix((eq_rows, total), dtype=float)
    b_eq = np.zeros(eq_rows)
    for t in range(n):
        a_eq[t, x0 + t] = 1.0
        a_eq[t, c0 + t] = -1.0
        a_eq[t, d0 + t] = 1.0
        a_eq[t, w0 + t] = -1.0
        b_eq[t] = load[t] - pv[t]

        row = n + t
        a_eq[row, s0 + t] = 1.0
        if t:
            a_eq[row, s0 + t - 1] = -1.0
            b_eq[row] = 0.0
        else:
            b_eq[row] = start_soc
        a_eq[row, c0 + t] = -eta
        a_eq[row, d0 + t] = 1.0 / eta
    next_row = 2 * n
    if adjusted:
        for t in range(n):
            a_eq[next_row + t, x0 + t] = 1.0
            a_eq[next_row + t, up0 + t] = -1.0
            a_eq[next_row + t, down0 + t] = 1.0
            b_eq[next_row + t] = float(old_purchase[t])
        next_row += n
    if terminal_soc is not None:
        a_eq[next_row, s0 + n - 1] = 1.0
        b_eq[next_row] = terminal_soc

    bounds = (
        [(0.0, None)] * n
        + [(0.0, FLOW_MAX)] * n
        + [(0.0, FLOW_MAX)] * n
        + [(SOC_MIN, SOC_MAX)] * n
        + [(0.0, None)] * n
    )
    if adjusted:
        bounds += [(0.0, None)] * (2 * n)
    result = linprog(objective, A_eq=a_eq.tocsr(), b_eq=b_eq, bounds=bounds, method="highs")
    if result.status != 0:
        raise RuntimeError(f"线性规划失败: {result.message}")
    vector = result.x
    purchase = vector[x0:x0 + n]
    charge = vector[c0:c0 + n]
    discharge = vector[d0:d0 + n]
    soc = vector[s0:s0 + n]
    curtail = vector[w0:w0 + n]
    residual = purchase + pv + discharge - load - charge - curtail
    return Plan(purchase, charge, discharge, soc, curtail, float(result.fun), float(np.max(np.abs(residual))))


def execute_committed(
    purchase: np.ndarray,
    actual_load: np.ndarray,
    actual_pv: np.ndarray,
    start_soc: float,
) -> Execution:
    n = len(purchase)
    charge = np.zeros(n)
    discharge = np.zeros(n)
    emergency = np.zeros(n)
    curtail = np.zeros(n)
    soc = np.zeros(n)
    current = float(start_soc)
    for t in range(n):
        surplus = purchase[t] + actual_pv[t] - actual_load[t]
        if surplus >= 0:
            charge[t] = min(surplus, FLOW_MAX, (SOC_MAX - current) / ETA)
            curtail[t] = surplus - charge[t]
        else:
            deficit = -surplus
            discharge[t] = min(deficit, FLOW_MAX, (current - SOC_MIN) * ETA)
            emergency[t] = deficit - discharge[t]
        current += ETA * charge[t] - discharge[t] / ETA
        soc[t] = current
    residual = purchase + actual_pv + discharge + emergency - actual_load - charge - curtail
    return Execution(charge, discharge, soc, emergency, curtail, float(np.max(np.abs(residual))))


def execute_segment(
    purchase: np.ndarray,
    actual_load: np.ndarray,
    actual_pv: np.ndarray,
    start_soc: float,
) -> Execution:
    return execute_committed(purchase, actual_load, actual_pv, start_soc)


def interval_label(start_index: int, end_index: int | None = None) -> str:
    end_index = start_index + 1 if end_index is None else end_index
    start_minutes = start_index * 10
    end_minutes = end_index * 10
    sh, sm = divmod(start_minutes, 60)
    eh, em = divmod(end_minutes, 60)
    return f"{sh:02d}:{sm:02d}-{eh:02d}:{em:02d}"


def emergency_runs(values: np.ndarray, threshold: float = 1e-7) -> list[tuple[int, int, float]]:
    runs: list[tuple[int, int, float]] = []
    active = np.flatnonzero(values > threshold)
    if not len(active):
        return runs
    start = previous = int(active[0])
    for index in active[1:]:
        index = int(index)
        if index != previous + 1:
            runs.append((start, previous + 1, float(values[start:previous + 1].sum())))
            start = index
        previous = index
    runs.append((start, previous + 1, float(values[start:previous + 1].sum())))
    return runs


def date_index(dates: np.ndarray, date_string: str) -> int:
    matches = np.flatnonzero(dates == np.datetime64(date_string))
    if len(matches) != 1:
        raise KeyError(date_string)
    return int(matches[0])
