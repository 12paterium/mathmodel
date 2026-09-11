"""Independently validate generated workbooks and machine-readable result arrays."""

from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import openpyxl

from model_core import ATTACHMENTS, ETA, FLOW_MAX, N_INTERVALS, ROOT, SOC_MAX, SOC_MIN, load_inputs


RESULTS = ROOT / "results"


def check_close(name: str, value: float, tolerance: float) -> None:
    if abs(value) > tolerance:
        raise AssertionError(f"{name}: {value} > {tolerance}")


def validate_problem1(inputs) -> dict[str, float]:
    detail = np.load(RESULTS / "problem1_detail.npz")
    purchase = detail["purchase"]
    charge = detail["charge"]
    discharge = detail["discharge"]
    soc = detail["soc"]
    curtail = detail["curtail"]
    residual = purchase + inputs.fixed_pv + discharge - inputs.fixed_load - charge - curtail
    state = np.r_[6000.0, soc[:-1]] + ETA * charge - discharge / ETA - soc
    checks = {
        "balance": float(np.max(np.abs(residual))),
        "state": float(np.max(np.abs(state))),
        "end_soc": float(abs(soc[-1] - 6000.0)),
        "flow": float(max(charge.max(), discharge.max()) - FLOW_MAX),
    }
    check_close("problem1 balance", checks["balance"], 1e-6)
    check_close("problem1 state", checks["state"], 1e-6)
    check_close("problem1 end SOC", checks["end_soc"], 1e-6)
    if soc.min() < SOC_MIN - 1e-6 or soc.max() > SOC_MAX + 1e-6:
        raise AssertionError("problem1 SOC越界")
    return checks


def validate_scenario(name: str, inputs) -> dict[str, float]:
    detail = np.load(RESULTS / f"{name}_detail.npz")
    purchase = detail["final_purchase"]
    charge = detail["charge"]
    discharge = detail["discharge"]
    soc = detail["soc"]
    emergency = detail["emergency"]
    curtail = detail["curtail"]
    load = inputs.load[31:]
    pv = inputs.pv[31:]
    residual = purchase + pv + discharge + emergency - load - charge - curtail
    previous = np.empty_like(soc)
    previous[:, 0] = detail["start_soc"]
    previous[:, 1:] = soc[:, :-1]
    state = previous + ETA * charge - discharge / ETA - soc
    continuity = detail["start_soc"][1:] - soc[:-1, -1]
    checks = {
        "balance": float(np.max(np.abs(residual))),
        "state": float(np.max(np.abs(state))),
        "continuity": float(np.max(np.abs(continuity))),
        "soc_lower": float(max(0.0, SOC_MIN - soc.min())),
        "soc_upper": float(max(0.0, soc.max() - SOC_MAX)),
        "flow": float(max(0.0, charge.max() - FLOW_MAX, discharge.max() - FLOW_MAX)),
        "simultaneous_charge_discharge": float(np.minimum(charge, discharge).max()),
    }
    for key, value in checks.items():
        check_close(f"{name} {key}", value, 1e-6)
    cost_sum = detail["plan_cost"] + detail["adjustment_cost"] + detail["emergency_cost"]
    check_close(f"{name} cost", float(np.max(np.abs(cost_sum - detail["total_cost"]))), 0.01)
    return checks


def validate_workbooks() -> dict[str, object]:
    expected = ["result1.xlsx", "result2.xlsx", "result3.xlsx", "result4-2.xlsx", "result4-3.xlsx"]
    result = {}
    for filename in expected:
        output = openpyxl.load_workbook(RESULTS / filename, read_only=True, data_only=True)
        template = openpyxl.load_workbook(ATTACHMENTS / "附件5" / filename, read_only=True, data_only=True)
        if output.sheetnames != template.sheetnames:
            raise AssertionError(f"{filename}工作表名称或顺序改变")
        first = output.worksheets[0]
        if filename == "result1.xlsx":
            values = [first.cell(row, 2).value for row in range(2, 146)]
        else:
            values = [first.cell(2, column).value for column in range(2, 146)]
        if len(values) != N_INTERVALS or any(value is None for value in values):
            raise AssertionError(f"{filename}未完整填充144个区间")
        result[filename] = {"sheets": output.sheetnames, "size": (first.max_row, first.max_column)}
    return result


def main() -> None:
    inputs = load_inputs()
    checks = {"problem1": validate_problem1(inputs)}
    for name in ("problem2", "problem3", "problem4_2", "problem4_3", "problem4_2_perfect_price"):
        checks[name] = validate_scenario(name, inputs)
    checks["workbooks"] = validate_workbooks()
    checks["status"] = "PASS"
    (RESULTS / "validation.json").write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(checks, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
